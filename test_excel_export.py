#!/usr/bin/env python3
"""
Backend test suite for ProFinance Interior - Excel Export Feature
Tests: GET /api/projects/{project_id}/recap/xlsx endpoint
"""
import requests
import json
import time
import io
from datetime import datetime

# Configuration
BASE_URL = "https://join-interior-setup.preview.emergentagent.com/api"
PREMIUM_EMAIL = "furnitrue.mail@gmail.com"
PREMIUM_PASSWORD = "Password123"
EXISTING_PROJECT_ID = "c420b3e7-cd4a-4cad-93cb-e9385960d4f5"

# Test state
token = None
test_project_id = None
created_invoices = []

def log(msg):
    print(f"[{datetime.now().strftime('%H:%M:%S')}] {msg}")

def login_premium():
    """Login with premium account"""
    global token
    log("=== TEST 1: Login Premium Account ===")
    resp = requests.post(f"{BASE_URL}/auth/login", json={
        "email": PREMIUM_EMAIL,
        "password": PREMIUM_PASSWORD
    })
    assert resp.status_code == 200, f"Login failed: {resp.status_code} {resp.text}"
    data = resp.json()
    token = data["token"]
    log(f"✓ Login successful, token: {token[:20]}...")
    return token

def setup_test_data():
    """Setup test data: Create invoices with variations"""
    global test_project_id
    log("\n=== TEST 2: Setup Test Data - Create Invoices ===")
    headers = {"Authorization": f"Bearer {token}"}
    
    # Use existing project
    test_project_id = EXISTING_PROJECT_ID
    log(f"Using existing project: {test_project_id}")
    
    # Invoice A: Proforma with quotationNo
    log("Creating Invoice A: Proforma, Terkirim, with quotationNo")
    inv_a_data = {
        "type": "proforma",
        "number": "XLS-PRO/001",
        "quotationNo": "QTO/2026/09/010",
        "status": "Terkirim",
        "clientName": "Bpk Uji Excel",
        "clientAddress": "Jl. Test Excel No. 1",
        "clientPhone": "081234567890",
        "items": [
            {
                "description": "DP 50%",
                "qty": 1,
                "unitPrice": 10000000
            }
        ],
        "ppnEnabled": True,
        "ppnPercent": 11,
        "retentionEnabled": False,
        "retentionPercent": 5,
        "companyName": "PT Interior Design",
        "companyAddress": "Jl. Company No. 1",
        "companyPhone": "021-12345678",
        "bankName": "BCA",
        "bankAccount": "1234567890",
        "bankHolder": "PT Interior Design"
    }
    
    resp_a = requests.post(
        f"{BASE_URL}/projects/{test_project_id}/invoices",
        headers=headers,
        json=inv_a_data
    )
    assert resp_a.status_code == 200, f"Create invoice A failed: {resp_a.status_code} {resp_a.text}"
    inv_a = resp_a.json()
    created_invoices.append(inv_a["id"])
    log(f"✓ Invoice A created: {inv_a['id']}, number={inv_a['number']}, quotationNo={inv_a['quotationNo']}")
    
    # Invoice B: Final with retention
    log("Creating Invoice B: Final, Lunas, with retention")
    inv_b_data = {
        "type": "final",
        "number": "XLS-INV/001",
        "quotationNo": "QTO/2026/09/010",
        "status": "Lunas",
        "clientName": "Bpk Uji Excel",
        "clientAddress": "Jl. Test Excel No. 1",
        "clientPhone": "081234567890",
        "items": [
            {
                "description": "Pelunasan",
                "qty": 1,
                "unitPrice": 20000000
            }
        ],
        "ppnEnabled": True,
        "ppnPercent": 11,
        "retentionEnabled": True,
        "retentionPercent": 5,
        "companyName": "PT Interior Design",
        "companyAddress": "Jl. Company No. 1",
        "companyPhone": "021-12345678",
        "bankName": "BCA",
        "bankAccount": "1234567890",
        "bankHolder": "PT Interior Design"
    }
    
    resp_b = requests.post(
        f"{BASE_URL}/projects/{test_project_id}/invoices",
        headers=headers,
        json=inv_b_data
    )
    assert resp_b.status_code == 200, f"Create invoice B failed: {resp_b.status_code} {resp_b.text}"
    inv_b = resp_b.json()
    created_invoices.append(inv_b["id"])
    log(f"✓ Invoice B created: {inv_b['id']}, number={inv_b['number']}, retention={inv_b['computed']['retentionAmount']:,}")
    
    # Invoice C: Proforma Draft without quotationNo
    log("Creating Invoice C: Proforma, Draft, no quotationNo")
    inv_c_data = {
        "type": "proforma",
        "number": "XLS-PRO/002",
        "status": "Draft",
        "clientName": "Bpk Uji Excel",
        "clientAddress": "Jl. Test Excel No. 1",
        "clientPhone": "081234567890",
        "items": [
            {
                "description": "Termin",
                "qty": 1,
                "unitPrice": 5000000
            }
        ],
        "ppnEnabled": False,
        "ppnPercent": 11,
        "retentionEnabled": False,
        "retentionPercent": 5,
        "companyName": "PT Interior Design",
        "companyAddress": "Jl. Company No. 1",
        "companyPhone": "021-12345678",
        "bankName": "BCA",
        "bankAccount": "1234567890",
        "bankHolder": "PT Interior Design"
    }
    
    resp_c = requests.post(
        f"{BASE_URL}/projects/{test_project_id}/invoices",
        headers=headers,
        json=inv_c_data
    )
    assert resp_c.status_code == 200, f"Create invoice C failed: {resp_c.status_code} {resp_c.text}"
    inv_c = resp_c.json()
    created_invoices.append(inv_c["id"])
    log(f"✓ Invoice C created: {inv_c['id']}, number={inv_c['number']}, status={inv_c['status']}")
    
    log(f"\n✓ All 3 test invoices created successfully")
    return inv_a, inv_b, inv_c

def test_excel_export():
    """Test 3: GET /api/projects/{project_id}/recap/xlsx?auth={token}"""
    log("\n=== TEST 3: Excel Export - Main Test ===")
    
    resp = requests.get(f"{BASE_URL}/projects/{test_project_id}/recap/xlsx?auth={token}")
    
    # Verify HTTP 200
    assert resp.status_code == 200, f"Excel export failed: {resp.status_code} {resp.text}"
    log(f"✓ HTTP 200 OK")
    
    # Verify Content-Type
    content_type = resp.headers.get("Content-Type")
    expected_content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert content_type == expected_content_type, f"Wrong Content-Type: expected '{expected_content_type}', got '{content_type}'"
    log(f"✓ Content-Type correct: {content_type}")
    
    # Verify body not empty
    xlsx_bytes = resp.content
    assert len(xlsx_bytes) > 0, "Excel file is empty"
    log(f"✓ Body not empty: {len(xlsx_bytes)} bytes")
    
    # Verify ZIP magic bytes "PK"
    assert xlsx_bytes[:2] == b"PK", f"Not a valid ZIP/XLSX file (magic bytes: {xlsx_bytes[:4]})"
    log(f"✓ ZIP magic bytes 'PK' verified")
    
    # Open with openpyxl
    try:
        from openpyxl import load_workbook
    except ImportError:
        log("⚠ openpyxl not installed, installing...")
        import subprocess
        subprocess.check_call(["pip", "install", "-q", "openpyxl"])
        from openpyxl import load_workbook
    
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = wb.sheetnames
    log(f"✓ Opened with openpyxl, found {len(sheet_names)} sheets: {sheet_names}")
    
    # Verify required sheets
    assert "Rekap Penagihan" in sheet_names, f"Missing 'Rekap Penagihan' sheet. Found: {sheet_names}"
    log(f"✓ Sheet 'Rekap Penagihan' found")
    
    assert "Daftar Invoice" in sheet_names, f"Missing 'Daftar Invoice' sheet. Found: {sheet_names}"
    log(f"✓ Sheet 'Daftar Invoice' found")
    
    # Check for RAB sheet (optional, depends on project)
    if "RAB" in sheet_names:
        log(f"✓ Sheet 'RAB' found (project has RAB)")
    else:
        log(f"ℹ Sheet 'RAB' not found (project may not have RAB, this is OK)")
    
    # Verify "Daftar Invoice" sheet content
    log("\n--- Verifying 'Daftar Invoice' sheet content ---")
    ws_invoice = wb["Daftar Invoice"]
    
    # Find header row and data
    found_invoice_number = False
    found_quotation_header = False
    found_total_tagihan_header = False
    
    for row in ws_invoice.iter_rows(values_only=True):
        row_str = " ".join([str(cell) if cell is not None else "" for cell in row])
        
        # Check for our test invoice number
        if "XLS-PRO/001" in row_str:
            found_invoice_number = True
            log(f"✓ Found invoice 'XLS-PRO/001' in data row")
        
        # Check for headers
        if "No. Quotation" in row_str or "No Quotation" in row_str or "Quotation" in row_str:
            found_quotation_header = True
        if "Total Tagihan" in row_str:
            found_total_tagihan_header = True
    
    assert found_invoice_number, "Invoice 'XLS-PRO/001' not found in 'Daftar Invoice' sheet"
    assert found_quotation_header, "Header 'No. Quotation' not found in 'Daftar Invoice' sheet"
    log(f"✓ Header 'No. Quotation' found")
    assert found_total_tagihan_header, "Header 'Total Tagihan' not found in 'Daftar Invoice' sheet"
    log(f"✓ Header 'Total Tagihan' found")
    
    # Verify "Rekap Penagihan" sheet content
    log("\n--- Verifying 'Rekap Penagihan' sheet content ---")
    ws_recap = wb["Rekap Penagihan"]
    
    found_recap_title = False
    found_total_tertagih = False
    found_retensi_ditahan = False
    
    for row in ws_recap.iter_rows(values_only=True):
        row_str = " ".join([str(cell) if cell is not None else "" for cell in row])
        
        if "REKAP PENAGIHAN PROYEK" in row_str.upper():
            found_recap_title = True
            log(f"✓ Found 'REKAP PENAGIHAN PROYEK' title")
        
        if "Total Tertagih" in row_str or "TOTAL TERTAGIH" in row_str.upper():
            found_total_tertagih = True
            log(f"✓ Found 'Total Tertagih' label")
        
        if "Retensi Ditahan" in row_str or "RETENSI DITAHAN" in row_str.upper():
            found_retensi_ditahan = True
            log(f"✓ Found 'Retensi Ditahan' label")
    
    assert found_recap_title, "'REKAP PENAGIHAN PROYEK' title not found in 'Rekap Penagihan' sheet"
    assert found_total_tertagih, "'Total Tertagih' label not found in 'Rekap Penagihan' sheet"
    assert found_retensi_ditahan, "'Retensi Ditahan' label not found in 'Rekap Penagihan' sheet"
    
    wb.close()
    log(f"\n✓ All Excel content verification passed")
    
    return xlsx_bytes

def test_auth_no_token():
    """Test 4: GET /api/projects/{project_id}/recap/xlsx without auth"""
    log("\n=== TEST 4: Excel Export - No Auth (expect 401) ===")
    
    # Try without auth header and without query param
    resp = requests.get(f"{BASE_URL}/projects/{test_project_id}/recap/xlsx")
    
    assert resp.status_code == 401, f"Expected 401 without auth, got {resp.status_code}"
    log(f"✓ Correctly returned 401 without auth")

def test_free_user_gating():
    """Test 5: Free user gating"""
    log("\n=== TEST 5: Excel Export - Free User Gating ===")
    
    # Register free user
    free_email = f"test_excel_free_{int(time.time())}@test.com"
    log(f"Registering free user: {free_email}")
    resp = requests.post(f"{BASE_URL}/auth/register", json={
        "email": free_email,
        "name": "Test Free User",
        "password": "Test123456"
    })
    assert resp.status_code == 200, f"Register failed: {resp.status_code} {resp.text}"
    free_token = resp.json()["token"]
    log(f"✓ Free user registered")
    
    # Try to access premium project with free token (via header)
    log("Testing free user accessing premium project...")
    headers = {"Authorization": f"Bearer {free_token}"}
    resp = requests.get(f"{BASE_URL}/projects/{test_project_id}/recap/xlsx", headers=headers)
    
    # Expected: 403 (premium required) or 404 (not owned)
    log(f"  Response: {resp.status_code}")
    assert resp.status_code in [403, 404], f"Expected 403 or 404 for free user accessing premium project, got {resp.status_code}"
    
    if resp.status_code == 403:
        log(f"✓ Free user correctly blocked with 403 (require_premium)")
    else:
        log(f"✓ Free user correctly blocked with 404 (not owned)")
    
    # Create a project for free user
    log("Creating project for free user...")
    resp = requests.post(f"{BASE_URL}/projects", headers=headers, json={
        "name": "Free User Test Project",
        "owner": "Bpk Free",
        "nominal": 10000000,
        "category": "Residensial",
        "status": "Berjalan"
    })
    assert resp.status_code == 200, f"Create project failed: {resp.status_code} {resp.text}"
    free_project_id = resp.json()["id"]
    log(f"✓ Free user project created: {free_project_id}")
    
    # Try to export free user's own project
    log("Testing free user exporting own project...")
    resp = requests.get(f"{BASE_URL}/projects/{free_project_id}/recap/xlsx", headers=headers)
    
    # Expected: 403 (require_premium)
    assert resp.status_code == 403, f"Expected 403 (require_premium) for free user exporting own project, got {resp.status_code}"
    log(f"✓ Free user correctly blocked with 403 (require_premium) even for own project")
    
    # Cleanup free user project
    requests.delete(f"{BASE_URL}/projects/{free_project_id}", headers=headers)

def cleanup():
    """Delete test invoices"""
    log("\n=== CLEANUP ===")
    headers = {"Authorization": f"Bearer {token}"}
    
    # Delete invoices
    for invoice_id in created_invoices:
        try:
            resp = requests.delete(f"{BASE_URL}/invoices/{invoice_id}", headers=headers)
            if resp.status_code == 200:
                log(f"✓ Deleted invoice: {invoice_id}")
            else:
                log(f"⚠ Failed to delete invoice {invoice_id}: {resp.status_code}")
        except Exception as e:
            log(f"⚠ Error deleting invoice {invoice_id}: {e}")

def main():
    """Run all tests"""
    print("=" * 80)
    print("ProFinance Interior Backend Test Suite")
    print("Testing: Excel Export Endpoint (GET /api/projects/{id}/recap/xlsx)")
    print("=" * 80)
    
    try:
        # Test 1: Login
        login_premium()
        
        # Test 2: Setup test data
        setup_test_data()
        
        # Test 3: Excel export main test
        test_excel_export()
        
        # Test 4: Auth test
        test_auth_no_token()
        
        # Test 5: Free user gating
        test_free_user_gating()
        
        # Cleanup
        cleanup()
        
        print("\n" + "=" * 80)
        print("✓ ALL TESTS PASSED (5/5)")
        print("=" * 80)
        print("\nSUMMARY:")
        print("✅ TEST 1: Login premium account - PASSED")
        print("✅ TEST 2: Setup test data (3 invoices) - PASSED")
        print("✅ TEST 3: Excel export (200, correct Content-Type, ZIP magic bytes, sheets verified) - PASSED")
        print("✅ TEST 4: Auth test (401 without token) - PASSED")
        print("✅ TEST 5: Free user gating (403 require_premium) - PASSED")
        
    except AssertionError as e:
        print(f"\n❌ TEST FAILED: {e}")
        cleanup()
        exit(1)
    except Exception as e:
        print(f"\n❌ ERROR: {e}")
        import traceback
        traceback.print_exc()
        cleanup()
        exit(1)

if __name__ == "__main__":
    main()
